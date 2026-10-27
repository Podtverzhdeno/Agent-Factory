"""Synthetic-only two-task B2B demo: XLSX RFQ and read-only supplier check."""
from __future__ import annotations
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from agent_builder import Builder, BuildCommand, Limits, TaskSnapshot, ToolManifest
from .core import Decision, EffectLedger, Runtime
from .verified_tools import ToolPolicy, VerifiedToolGateway

SHA = 'sha256:' + 'b' * 64

class DemoRegistry:
    def eligible_manifests(self, tenant_id):
        rows = [('read_xlsx','source.xlsx','read','files:read'),
                ('catalog_lookup','catalog.lookup','read','catalog:read'),
                ('draft_xlsx','draft.xlsx','draft_file','draft:write'),
                ('supplier_check','supplier.lookup','read','supplier:read')]
        return [ToolManifest(tool_id=tool, version='1.0.0', manifest_digest=SHA,
            capabilities=frozenset([cap]), side_effects=frozenset([effect]),
            required_scopes=frozenset([scope]), tenant_availability=frozenset([tenant_id]),
            input_schema_ref='artifact://schemas/demo-input',
            output_schema_ref='artifact://schemas/demo-output',
            risk_level='low', approved=True, transport='LOCAL')
            for cap, tool, effect, scope in rows]

class DemoTools:
    def __init__(self, input_path: Path, output_dir: Path):
        self.input_path, self.output_dir = input_path, output_dir
        self.calls = []
    async def call(self, *, tenant_id, tool_id, args, idempotency_key):
        self.calls.append((tenant_id, tool_id, idempotency_key))
        if tool_id == 'source.xlsx':
            book = load_workbook(self.input_path, data_only=True, read_only=True)
            try:
                rows = list(book.active.values)
            finally:
                book.close()
            return {'rows':[{'sku':str(row[0]),'qty':int(row[1])} for row in rows[1:]]}
        if tool_id == 'catalog.lookup':
            if args.get('sku') != 'BOLT-1':
                raise ValueError('unknown product')
            return {'sku':'BOLT-1','usd':12}
        if tool_id == 'supplier.lookup':
            return {'supplier':'SUP-1','status':'APPROVED',
                    'source':'artifact://synthetic/suppliers-v1'}
        if tool_id == 'draft.xlsx':
            self.output_dir.mkdir(parents=True, exist_ok=True)
            book = Workbook()
            ws = book.active
            ws.append(['SKU','Quantity','Unit USD','Total USD'])
            ws.append(['BOLT-1', args['qty'], args['unit_usd'], args['qty'] * args['unit_usd']])
            book.save(self.output_dir / 'draft-quote.xlsx')
            return {'artifact_ref':'artifact://synthetic/draft-quote.xlsx'}
        raise ValueError('unregistered tool')

class DemoModel:
    async def decide(self, *, route, prompt, history, tool_ids, trace_id):
        seen = [m for m in history if m['role'] == 'tool']
        if 'source.xlsx' in tool_ids:
            if len(seen) == 0:
                return Decision(kind='tool', tool_id='source.xlsx', args={})
            if len(seen) == 1:
                return Decision(kind='tool', tool_id='catalog.lookup', args={'sku':'BOLT-1'})
            if len(seen) == 2:
                qty = json.loads(seen[0]['content'])['rows'][0]['qty']
                unit = json.loads(seen[1]['content'])['usd']
                return Decision(kind='tool',tool_id='draft.xlsx',args={'qty':qty,'unit_usd':unit})
            return Decision(kind='final',text='Draft RFQ created with evidence')
        if not seen:
            return Decision(kind='tool',tool_id='supplier.lookup',args={'supplier':'SUP-1'})
        return Decision(kind='final',text='Supplier verified by synthetic source')

async def execute_two_tasks(work_dir: Path) -> tuple:
    work_dir.mkdir(parents=True,exist_ok=True)
    wb = Workbook()
    wb.active.append(['SKU','Quantity'])
    wb.active.append(['BOLT-1',2])
    input_path = work_dir / 'rfq.xlsx'
    wb.save(input_path)
    outcomes = []
    cases = [('rfq',('read_xlsx','catalog_lookup','draft_xlsx'),
              ('read','draft_file'),('files:read','catalog:read','draft:write')),
             ('supplier',('supplier_check',),('read',),('supplier:read',))]
    for name, capabilities, effects, scopes in cases:
        snapshot = TaskSnapshot(tenant_id='synthetic-tenant',task_id=name,
            task_snapshot_id='snapshot-' + name,status='READY',
            objective='Complete the synthetic B2B workflow ' + name,
            required_capabilities=capabilities,allowed_effects=frozenset(effects),
            authorized_scopes=frozenset(scopes),
            acceptance_policy_ref='artifact://policies/acceptance-v1',
            snapshot_digest=SHA,limits=Limits(max_steps=6,max_usd_per_run=1))
        command = BuildCommand(build_id='build-' + name,tenant_id=snapshot.tenant_id,
            task_snapshot_id=snapshot.task_snapshot_id,idempotency_key='build-' + name,
            capability_registry_revision='demo-v1',model_route='mock-model',
            evaluation_policy_ref=snapshot.acceptance_policy_ref)
        spec = Builder(DemoRegistry(),frozenset({'mock-model'})).assemble(snapshot,command)
        policies = {b.tool_id:ToolPolicy(tenant_id=spec.tenant_id,tool_id=b.tool_id,
            version=b.version,manifest_digest=b.digest,scopes=frozenset(b.granted_scopes),
            input_schema={'type':'object'},output_schema={'type':'object'},
            effect_mode=b.side_effects[0]) for b in spec.selected_tools}
        tools = DemoTools(input_path, work_dir)
        runtime = Runtime(DemoModel(), VerifiedToolGateway(tools,EffectLedger(),policies=policies),
            approved_spec_digests=frozenset({spec.artifact_digest}))
        outcome = await runtime.run(spec,tenant_id=spec.tenant_id,
            run_id='run-' + name,thread_id='thread-' + name,
            request='Run approved synthetic task',trace_id='trace-' + name)
        outcomes.append((spec,outcome,tools))
    return tuple(outcomes)
